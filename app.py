import os, re, json, unicodedata, difflib, sqlite3, threading, hashlib, secrets
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
app.config["SESSION_COOKIE_SECURE"]   = False
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_NAME"]     = "approw_session"

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))

# ── Stockage persistant : utilise /data si dispo (Railway Volume), sinon BASE_DIR ──
DATA_DIR = "/data" if os.path.isdir("/data") else BASE_DIR

MERE_PATH  = os.path.join(DATA_DIR, "mere.xlsx")
LOG_PATH   = os.path.join(DATA_DIR, "log.json")
USERS_PATH = os.path.join(DATA_DIR, "users.json")
UPLOAD_DIR = os.path.join(DATA_DIR, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

# ── Synchroniser mere.xlsx depuis BASE_DIR vers DATA_DIR si plus récent ou absent ──
_mere_src = os.path.join(BASE_DIR, "mere.xlsx")
if os.path.exists(_mere_src):
    import shutil as _shutil
    _should_copy = not os.path.exists(MERE_PATH)
    if not _should_copy:
        # Copier si le fichier source est plus grand (plus de données)
        _should_copy = os.path.getsize(_mere_src) > os.path.getsize(MERE_PATH)
    if _should_copy:
        _shutil.copy2(_mere_src, MERE_PATH)
        print(f"✅ mere.xlsx synchronisé vers {DATA_DIR} ({os.path.getsize(MERE_PATH)//1024} Ko)", flush=True)

# Admin password (set via Railway env var ADMIN_PASSWORD)
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin2026")

NCOLS  = 9
C_NEW  = "FFF2CC"
thin   = Side(style="thin", color="CCCCCC")
BDR    = Border(left=thin, right=thin, top=thin, bottom=thin)

write_lock = threading.Lock()
_db = None

def strip_accents(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()

def normalize(s):
    s = strip_accents(s)
    return re.sub(r"[-_\s\'\u2019\u2018]+", " ", s).strip()

def clean_phone(p):
    # Supprime espaces, tirets, points, parenthèses ET étoiles (*01..., *56...)
    p = re.sub(r"[\s\-\.\(\)\*]+", "", str(p or ""))
    return re.sub(r"^\+229|^00229", "", p).strip()

def load_log():
    if os.path.exists(LOG_PATH):
        with open(LOG_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {"sessions": [], "total_inserted": 0, "users": {}}

def save_log(log):
    with open(LOG_PATH, "w", encoding="utf-8") as f:
        json.dump(log, f, ensure_ascii=False, indent=2)

# ── User management ───────────────────────────────────────────────────────────

def hash_pw(pw):
    return hashlib.sha256(pw.encode("utf-8")).hexdigest()

def load_users():
    if os.path.exists(USERS_PATH):
        with open(USERS_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {}

def save_users(users):
    with open(USERS_PATH, "w", encoding="utf-8") as f:
        json.dump(users, f, ensure_ascii=False, indent=2)

def get_user(username):
    return load_users().get(username.lower().strip())

def create_user(username, password):
    users = load_users()
    key = username.lower().strip()
    if key in users:
        return False, "Utilisateur déjà existant"
    users[key] = {
        "display": username.strip(),
        "password": hash_pw(password),
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    save_users(users)
    return True, "Utilisateur créé"

def check_password(username, password):
    user = get_user(username)
    if not user: return False
    return user["password"] == hash_pw(password)

def is_admin():
    return session.get("role") == "admin"

def is_logged_in():
    return session.get("username") is not None or is_admin()

def current_username():
    if is_admin(): return "__admin__"
    return session.get("username", "")

def require_login(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not is_logged_in():
            return redirect(url_for("login_page"))
        return f(*args, **kwargs)
    return decorated

def require_admin(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not is_admin():
            return jsonify({"error": "Accès refusé"}), 403
        return f(*args, **kwargs)
    return decorated

def get_db():
    global _db
    if _db is None:
        _db = build_index_db()
    return _db

def invalidate_db():
    global _db
    _db = None

def build_index_db():
    print("\u26a1 Construction de l'index SQLite...", flush=True)
    t0 = datetime.now()
    wb = load_workbook(MERE_PATH, read_only=True)
    ws = wb.active
    con = sqlite3.connect(":memory:", check_same_thread=False)
    con.execute("""CREATE TABLE slots (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        village TEXT, village_norm TEXT,
        dept_norm TEXT, comm_norm TEXT, arrond_norm TEXT,
        row_start INTEGER, filled INTEGER, last_data_row INTEGER)""")
    con.execute("CREATE TABLE phones (phone TEXT PRIMARY KEY)")
    # arrond index: arrond_norm -> last_data_row + next_dept_row
    con.execute("""CREATE TABLE arrond_index (
        arrond_norm TEXT PRIMARY KEY,
        dept_norm TEXT, comm_norm TEXT,
        arrond_row INTEGER, hdr_row INTEGER,
        last_data_row INTEGER, next_dept_row INTEGER)""")

    dept_norm = comm_norm = arrond_norm = ""
    dept_row = arrond_row = hdr_row = 0
    arrond_data = {}   # arrond_norm -> dict
    slots = []
    phones = []

    all_rows = list(ws.iter_rows(values_only=True))
    total_rows = len(all_rows)

    for i, row in enumerate(all_rows):
        r  = i + 1
        v1 = str(row[0] or "").strip()
        v5 = str(row[4] or "").strip() if len(row) > 4 else ""
        v2 = str(row[1] or "").strip() if len(row) > 1 else ""
        v9 = str(row[8] or "").strip() if len(row) > 8 else ""

        if "DEPARTEMENT" in v1.upper():
            # Close previous arrond block
            if arrond_norm and arrond_norm in arrond_data:
                arrond_data[arrond_norm]["next_dept_row"] = r
            dept_norm = normalize(v1.replace("DEPARTEMENT:", "").strip())
            comm_norm = normalize(v5.replace("COMMUNE:", "").strip())

        elif "ARRONDISSEMENT" in v1.upper():
            arrond_norm = normalize(v1)
            arrond_row  = r
            arrond_data[arrond_norm] = {
                "dept_norm": dept_norm, "comm_norm": comm_norm,
                "arrond_row": r, "hdr_row": 0,
                "last_data_row": r, "next_dept_row": total_rows + 1
            }

        elif v1 == "N\u00b0" and arrond_norm:
            arrond_data[arrond_norm]["hdr_row"] = r

        elif v1 == "1" and v2 and arrond_norm:
            slots.append((v2, normalize(v2), dept_norm, comm_norm, arrond_norm, r, 0, r + 4))

        if v1 in ["1","2","3","4","5"] and arrond_norm:
            arrond_data[arrond_norm]["last_data_row"] = r

        if v9:
            p = clean_phone(v9)
            if p: phones.append((p,))

    wb.close()

    # Calculer filled directement depuis all_rows (déjà en mémoire — pas besoin de relire)
    row_to_slot_idx  = {}  # row_start -> index dans slots[]
    village_to_slot_idx = {}  # village_norm -> index (pour overflow N>5)
    for idx, s in enumerate(slots):
        row_to_slot_idx[s[5]] = idx  # s[5] = row_start
        if s[1] not in village_to_slot_idx:   # s[1] = village_norm
            village_to_slot_idx[s[1]] = idx

    filled_counts  = [0] * len(slots)
    last_data_rows = [s[7] for s in slots]  # s[7] = last_data_row initial

    for i, row in enumerate(all_rows):
        r  = i + 1
        v1 = str(row[0] or "").strip()
        v2 = str(row[1] or "").strip() if len(row) > 1 else ""
        v3 = str(row[2] or "").strip() if len(row) > 2 else ""
        if not v1.isdigit() or not v3: continue
        n = int(v1)
        if n >= 1 and n <= 5:
            # Ligne normale (N=1-5) → mise à jour filled + last_data_row
            slot_row = r - (n - 1)
            if slot_row in row_to_slot_idx:
                idx = row_to_slot_idx[slot_row]
                filled_counts[idx]  = max(filled_counts[idx], n)
                last_data_rows[idx] = max(last_data_rows[idx], r)
        elif n > 5 and v2:
            # Ligne overflow (N>5) → mise à jour filled ET last_data_row
            v2_norm = normalize(v2)
            if v2_norm in village_to_slot_idx:
                idx = village_to_slot_idx[v2_norm]
                filled_counts[idx]  = max(filled_counts[idx], n)  # ← compter overflow aussi
                last_data_rows[idx] = max(last_data_rows[idx], r)

    # Appliquer les filled counts
    slots_with_filled = [
        (s[0], s[1], s[2], s[3], s[4], s[5], filled_counts[i], last_data_rows[i])
        for i, s in enumerate(slots)
    ]

    con.executemany(
        "INSERT INTO slots(village,village_norm,dept_norm,comm_norm,arrond_norm,row_start,filled,last_data_row) VALUES(?,?,?,?,?,?,?,?)",
        slots_with_filled)
    con.executemany("INSERT OR IGNORE INTO phones VALUES(?)", phones)
    con.executemany(
        "INSERT OR REPLACE INTO arrond_index VALUES(?,?,?,?,?,?,?)",
        [(k, v["dept_norm"], v["comm_norm"], v["arrond_row"], v["hdr_row"],
          v["last_data_row"], v["next_dept_row"])
         for k, v in arrond_data.items()])

    con.execute("CREATE INDEX idx_vn  ON slots(village_norm)")
    con.execute("CREATE INDEX idx_an  ON slots(arrond_norm)")
    con.execute("CREATE INDEX idx_van ON slots(village_norm, arrond_norm)")
    con.commit()

    elapsed = (datetime.now() - t0).total_seconds()
    n  = con.execute("SELECT COUNT(*) FROM slots").fetchone()[0]
    np = con.execute("SELECT COUNT(*) FROM phones").fetchone()[0]
    na = con.execute("SELECT COUNT(*) FROM arrond_index").fetchone()[0]
    print(f"   \u2705 Index pr\u00eat en {elapsed:.1f}s \u2014 {n} villages, {na} arrond., {np} t\u00e9l.", flush=True)
    return con

def find_slot(con, quartier, arrond_norm_ctx=None, dept_norm_ctx=None, comm_norm_ctx=None):
    """
    Cherche un slot par quartier avec contexte géographique précis.
    Priorité : dept+comm+arrond > arrond seul > global
    Cela évite de confondre deux quartiers homonymes dans des arronds différents.
    """
    key = normalize(quartier)

    # ── Niveau 1 : Dept + Commune + Arrond (le plus précis) ──────────────────
    if dept_norm_ctx and comm_norm_ctx and arrond_norm_ctx:
        row = con.execute(
            "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
            "WHERE village_norm=? AND dept_norm=? AND comm_norm=? AND arrond_norm=? LIMIT 1",
            (key, dept_norm_ctx, comm_norm_ctx, arrond_norm_ctx)).fetchone()
        if row: return row
        # Fuzzy dans dept+comm+arrond
        keys = [r[0] for r in con.execute(
            "SELECT DISTINCT village_norm FROM slots "
            "WHERE dept_norm=? AND comm_norm=? AND arrond_norm=?",
            (dept_norm_ctx, comm_norm_ctx, arrond_norm_ctx)).fetchall()]
        matches = difflib.get_close_matches(key, keys, n=1, cutoff=0.85)
        if matches:
            row = con.execute(
                "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
                "WHERE village_norm=? AND dept_norm=? AND comm_norm=? AND arrond_norm=? LIMIT 1",
                (matches[0], dept_norm_ctx, comm_norm_ctx, arrond_norm_ctx)).fetchone()
            if row: return row

    # ── Niveau 2 : Arrond seul ────────────────────────────────────────────────
    if arrond_norm_ctx:
        row = con.execute(
            "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
            "WHERE village_norm=? AND arrond_norm=? LIMIT 1",
            (key, arrond_norm_ctx)).fetchone()
        if row: return row
        keys = [r[0] for r in con.execute(
            "SELECT DISTINCT village_norm FROM slots WHERE arrond_norm=?",
            (arrond_norm_ctx,)).fetchall()]
        matches = difflib.get_close_matches(key, keys, n=1, cutoff=0.85)
        if matches:
            row = con.execute(
                "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
                "WHERE village_norm=? AND arrond_norm=? LIMIT 1",
                (matches[0], arrond_norm_ctx)).fetchone()
            if row: return row

    # ── Niveau 3 : Global (fallback si aucun contexte) ───────────────────────
    row = con.execute(
        "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
        "WHERE village_norm=? LIMIT 1", (key,)).fetchone()
    if row: return row
    all_keys = [r[0] for r in con.execute("SELECT DISTINCT village_norm FROM slots").fetchall()]
    matches = difflib.get_close_matches(key, all_keys, n=1, cutoff=0.85)
    if matches:
        return con.execute(
            "SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots "
            "WHERE village_norm=? LIMIT 1", (matches[0],)).fetchone()
    return None


def find_arrond_in_mere(con, arrond_name):
    """Cherche l'arrondissement dans le fichier mère par fuzzy matching."""
    key = normalize(arrond_name)
    # Exact
    row = con.execute(
        "SELECT arrond_norm, last_data_row, next_dept_row, hdr_row FROM arrond_index "
        "WHERE arrond_norm=? LIMIT 1", (key,)).fetchone()
    if row: return row
    # Fuzzy
    all_keys = [r[0] for r in con.execute("SELECT arrond_norm FROM arrond_index").fetchall()]
    matches = difflib.get_close_matches(key, all_keys, n=1, cutoff=0.80)
    if matches:
        return con.execute(
            "SELECT arrond_norm, last_data_row, next_dept_row, hdr_row FROM arrond_index "
            "WHERE arrond_norm=? LIMIT 1", (matches[0],)).fetchone()
    return None

def style_new_row(ws, rn):
    ws.row_dimensions[rn].height = 15
    for c in range(1, NCOLS+1):
        cell = ws.cell(row=rn, column=c)
        cell.fill      = PatternFill("solid", start_color=C_NEW)
        cell.font      = Font(name="Arial", size=10)
        cell.border    = BDR
        cell.alignment = Alignment(vertical="center")
    ws.cell(row=rn, column=1).alignment = Alignment(horizontal="center", vertical="center")

def write_person_row(ws, rn, num, village, p):
    ws.cell(row=rn, column=1).value = num
    ws.cell(row=rn, column=2).value = village
    ws.cell(row=rn, column=3).value = p["nom"]
    ws.cell(row=rn, column=4).value = p["prenom"]
    ws.cell(row=rn, column=5).value = p.get("partis","")
    ws.cell(row=rn, column=6).value = p.get("profession","")
    ws.cell(row=rn, column=7).value = p.get("date_naissance","")
    ws.cell(row=rn, column=8).value = p.get("lieu_naissance","")
    ws.cell(row=rn, column=9).value = p["telephone"]

ALIASES = {
    "quartier":       ["quartier","village","localite","quartier/village","quartiers",
                       "villages / quartiers de ville"],
    "nom":            ["nom","noms","name"],
    "prenom":         ["prenom","prenoms","firstname"],
    "nom_prenom":     ["nom et prenoms","nom et prenom","noms et prenoms","responsable"],
    "telephone":      ["telephone","tel","phone","mobile","contact",
                       "numero de telephone","numero",
                       "adresse complete","adresse complete","adresse",
                       "adresse compete"],
    "partis":         ["partis","parti","party"],
    "profession":     ["profession","metier","emploi"],
    "date_naissance": ["date de naissance","date_naissance","naissance"],
    "lieu_naissance": ["lieu de naissance","lieu_naissance"],
}

def split_nom_prenom(full_name):
    """
    Premier mot = NOM, tout le reste = PRENOMS.
    Ex: 'LAWANI SENI Akim' → NOM='LAWANI', PRENOM='SENI Akim'
    Ex: 'GNAMMI Datonga'   → NOM='GNAMMI', PRENOM='Datonga'
    Si un seul mot : NOM=PRENOM=ce mot (évite le rejet).
    """
    parts = full_name.strip().split()
    if not parts:       return '', ''
    if len(parts) == 1: return parts[0], parts[0]
    return parts[0], ' '.join(parts[1:])

def detect_columns(ws_in, max_col=30):
    """
    Détecte les colonnes en scannant les premières lignes.
    Supporte aussi la colonne 'Nom et Prénoms' combinée.
    max_col : limite le scan pour éviter les fichiers avec 16384 colonnes vides.
    """
    sample_rows = list(ws_in.iter_rows(max_row=10, max_col=max_col, values_only=True))
    for r_idx, sample_row in enumerate(sample_rows):
        r = r_idx + 1
        mapping = {}
        for c_idx, val in enumerate(sample_row):
            c = c_idx + 1
            h = normalize(str(val or ""))
            for field, aliases in ALIASES.items():
                if h in aliases:
                    mapping[field] = c
        # Cas normal : nom + prenom séparés
        if "nom" in mapping and "prenom" in mapping:
            return mapping, r
        # Cas spécial : colonne combinée "Nom et Prénoms" ou "RESPONSABLE"
        if "nom_prenom" in mapping:
            mapping["nom"]    = mapping["nom_prenom"]
            mapping["prenom"] = mapping["nom_prenom"]
            mapping["_split_nom_prenom"] = True
            return mapping, r
    return {}, None

def integrate_file(filepath, filename, username=None):
    # ── Support .xls : convertir en xlsx à la volée ───────────────────────────
    if filepath.lower().endswith(".xls"):
        try:
            import xlrd
            xls_wb  = xlrd.open_workbook(filepath)
            xls_ws  = xls_wb.sheet_by_index(0)
            from openpyxl import Workbook as WB
            new_wb  = WB()
            new_ws  = new_wb.active
            for rr in range(xls_ws.nrows):
                row_vals = []
                for cc in range(xls_ws.ncols):
                    cell = xls_ws.cell(rr, cc)
                    # xlrd type 2 = float, might be phone number
                    if cell.ctype == 2 and cell.value == int(cell.value):
                        row_vals.append(str(int(cell.value)))
                    else:
                        row_vals.append(cell.value if cell.value != '' else None)
                new_ws.append(row_vals)
            xlsx_path = filepath + ".xlsx"
            new_wb.save(xlsx_path)
            filepath = xlsx_path
        except ImportError:
            return {"filename": filename,
                    "error": "Fichier .xls non supporté sur ce serveur. Veuillez le convertir en .xlsx avec Excel ou LibreOffice."}
        except Exception as e:
            return {"filename": filename, "error": f"Erreur lecture .xls : {e}"}

    # ── Ouverture en mode lecture seule (x70 plus rapide pour les gros fichiers) ──
    wb_in = load_workbook(filepath, data_only=True, read_only=True)
    ws_in = wb_in.active

    # Détecter le vrai nombre de colonnes (certains fichiers ont 16384 cols vides)
    MAX_COL_FILE = 15
    for row in ws_in.iter_rows(max_row=10, values_only=True):
        for i, v in enumerate(row):
            if v is not None and i + 2 > MAX_COL_FILE:
                MAX_COL_FILE = min(i + 5, 30)

    # Mettre toutes les lignes en cache (évite les accès répétés sur le fichier)
    all_input_rows = list(ws_in.iter_rows(max_col=MAX_COL_FILE, values_only=True))
    wb_in.close()

    def _cell(r, c):
        """Accès cellule depuis le cache mémoire."""
        if r < 1 or r > len(all_input_rows): return ""
        row_data = all_input_rows[r - 1]
        if c < 1 or c > len(row_data): return ""
        return str(row_data[c - 1] or "").strip()

    # Créer un objet factice pour detect_columns qui utilise le cache
    class _WsProxy:
        def iter_rows(self, max_row=None, max_col=None, values_only=True):
            end = min(max_row, len(all_input_rows)) if max_row else len(all_input_rows)
            for row in all_input_rows[:end]:
                yield row[:max_col] if max_col else row
        @property
        def max_row(self): return len(all_input_rows)
        @property
        def max_column(self): return MAX_COL_FILE

    ws_proxy = _WsProxy()
    col_map, header_row = detect_columns(ws_proxy, max_col=MAX_COL_FILE)
    if not col_map:
        return {"filename": filename, "error": "Colonnes non detectees (Nom + Prenom requis)"}

    def gv(r, field):
        col = col_map.get(field)
        return str(_cell(r, col)) if col else ""

    is_split = col_map.get("_split_nom_prenom", False)

    def find_tel(r, base_col):
        if not base_col: return ""
        for delta in [0, 1, -1, 2]:
            v = str(_cell(r, base_col+delta))
            if v and re.search(r"\d{5,}", v.replace(" ", "").replace("-","")):
                return v
        return ""

    # Détecter alt_map (colonnes décalées de +1) seulement si pas de mode split
    alt_map = None
    if not is_split:
        alt_hits = 0
        nom_col = col_map.get("nom", 99)
        if nom_col and nom_col <= MAX_COL_FILE:
            for r_idx in range(header_row, min(header_row + 60, len(all_input_rows))):
                nom_val  = _cell(r_idx + 1, nom_col)
                nom1_val = _cell(r_idx + 1, nom_col + 1)
                if not nom_val and nom1_val:
                    alt_hits += 1
        if alt_hits > 0:
            alt_map = {f: c+1 for f, c in col_map.items() if not f.startswith("_") and c < MAX_COL_FILE}

    def get_row_data(r):
        # ── Quartier ──────────────────────────────────────────────────────────
        quartier = gv(r, "quartier")
        if not quartier:
            col_nom_val = col_map.get("nom", col_map.get("nom_prenom", 99))
            skip = {"N","NOM","NOMS","PRENOM","PRENOMS","PROFESSION",
                    "ADRESSE COMPLETE","ADRESSE COMPETE","CONTACT",
                    "QUARTIER","QUARTIERS","N° QUARTIERS",""}

            # Cas 7ème CE : quartier en col_nom-1 (col juste avant Nom et Prénoms)
            col_before = col_nom_val - 1
            if col_before >= 1:
                v = str(_cell(r, col_before))
                norm_v = normalize(v).upper()
                if (v and norm_v not in [normalize(s) for s in skip]
                        and not re.match(r"^\d+$", v)
                        and len(v) > 1):
                    quartier = v

            # Cas général : chercher aussi col_quartier±1
            if not quartier and "quartier" in col_map:
                col_q = col_map["quartier"]
                for delta in [1, -1]:
                    v = str(_cell(r, col_q+delta))
                    if v and normalize(v) not in [normalize(s) for s in skip] \
                            and not re.match(r"^\d+$", v):
                        quartier = v; break

        # ── Nom / Prénom ──────────────────────────────────────────────────────
        if is_split:
            # Mode "Nom et Prénoms" combiné ou "RESPONSABLE" : splitter le contenu
            full = gv(r, "nom").strip()
            if full:
                nom, prenom = split_nom_prenom(full)
            else:
                nom, prenom = "", ""
        else:
            nom    = gv(r, "nom").strip()
            prenom = gv(r, "prenom").strip()

            # Cas RoW 4ème VF : NOM vide mais PRENOM contient nom+prénom combinés
            if not nom and prenom:
                nom, prenom = split_nom_prenom(prenom)

            # Cas RoW 4ème VF : NOM contient nom+prénom combinés, PRENOM vide
            elif nom and not prenom and " " in nom:
                # Vérifier si c'est vraiment un nom combiné (plusieurs mots)
                col_n = col_map.get("nom")
                col_p = col_map.get("prenom")
                if col_n and col_p and col_n != col_p:
                    nom, prenom = split_nom_prenom(nom)

        # ── Téléphone ─────────────────────────────────────────────────────────
        tel_raw = find_tel(r, col_map.get("telephone"))

        # Alt_map : colonnes décalées si ligne principale vide
        if alt_map and not nom:
            n2 = str(_cell(r, alt_map.get("nom",99)))
            if n2:
                q2 = str(_cell(r, alt_map.get("quartier",99)))
                p2 = str(_cell(r, alt_map.get("prenom",99)))
                if q2 and not quartier: quartier = q2
                if not p2 and " " in n2:
                    nom, prenom = split_nom_prenom(n2)
                else:
                    nom = n2; prenom = p2
                if not tel_raw: tel_raw = find_tel(r, alt_map.get("telephone"))

        return quartier, nom, prenom, tel_raw, {
            "partis":         gv(r, "partis"),
            "profession":     gv(r, "profession"),
            "date_naissance": gv(r, "date_naissance"),
            "lieu_naissance": gv(r, "lieu_naissance"),
        }

    persons = []; rejected = []; seen_phones = set()
    last_quartier = None
    last_arrond_ctx = None   # ← arrondissement courant du fichier source
    last_dept_ctx   = None   # ← département courant du fichier source
    last_comm_ctx   = None   # ← commune courante du fichier source

    for r in range(1, len(all_input_rows) + 1):
        # Détecter DEPARTEMENT / COMMUNE / ARRONDISSEMENT dans le fichier source
        for c in range(1, min(MAX_COL_FILE+1, 6)):
            cv = str(_cell(r, c))
            cv_up = cv.upper()
            if "DEPARTEMENT" in cv_up:
                last_dept_ctx = normalize(cv.replace("DEPARTEMENT:","").replace("DEPARTEMENT :","").strip())
            if "COMMUNE" in cv_up:
                last_comm_ctx = normalize(cv.replace("COMMUNE:","").replace("COMMUNE :","").strip())
            if "ARRONDISSEMENT" in cv_up:
                last_arrond_ctx = normalize(cv)
                break

        for c in range(1, min(MAX_COL_FILE+1, 6)):
            cv = str(_cell(r, c))
            m = re.match(r"^QUARTIER\s*[:\-]\s*(.+)", cv, re.IGNORECASE)
            if m: last_quartier = m.group(1).strip(); break
            if re.match(r"^QUARTIER\s*[:\-]?\s*$", cv, re.IGNORECASE):
                nxt = str(_cell(r, c+1))
                if nxt: last_quartier = nxt
                break
        if r < header_row+1: continue
        quartier, nom, prenom, tel_raw, extras = get_row_data(r)
        tel = clean_phone(tel_raw)
        if not any([quartier, nom, prenom, tel_raw]): continue
        if normalize(nom) in ["nom","noms","name"]: continue
        if normalize(prenom) in ["prenom","prenoms"]: continue
        first_cell = str(_cell(r, 1))
        if re.match(r"^QUARTIER\s*[:\-]", first_cell, re.IGNORECASE): continue
        skip_row = False
        for c in range(1, min(MAX_COL_FILE+1, 4)):
            cv = str(_cell(r, c)).upper()
            if cv.startswith("DEPARTEMENT") or "ARRONDISSEMENT" in cv or cv.startswith("COMMUNE"):
                skip_row = True; break
        if skip_row or not nom: continue
        if quartier: last_quartier = quartier
        elif not quartier and nom and (prenom or tel):
            if last_quartier: quartier = last_quartier
        missing = []
        if not quartier: missing.append("Quartier")
        if not nom:      missing.append("Nom")
        if not prenom:   missing.append("Prenom")
        if not tel:      missing.append("Telephone")
        if missing:
            reason_str = ", ".join(missing)
            rejected.append({"row": r, "quartier": quartier, "nom": nom,
                             "reason": "Manquant : " + reason_str})
            continue
        if tel in seen_phones:
            rejected.append({"row": r, "quartier": quartier, "nom": nom,
                             "reason": f"Doublon interne : {tel_raw}"}); continue
        seen_phones.add(tel)
        persons.append({"row": r, "quartier": quartier, "nom": nom, "prenom": prenom,
                        **extras, "telephone": tel_raw, "tel_clean": tel,
                        "arrond_ctx": last_arrond_ctx,
                        "dept_ctx":   last_dept_ctx,
                        "comm_ctx":   last_comm_ctx})

    _mere_max_row = 0  # sera défini dans le write_lock
    with write_lock:
        con    = get_db()
        wb_out = load_workbook(MERE_PATH)
        ws_out = wb_out.active
        _mere_max_row = ws_out.max_row
        inserted = overflow = new_quartier = duplicates = 0

        # ── Phase 1 : Résoudre tous les placements AVANT toute écriture ──────
        direct_writes  = []   # (target_row, person) pour slots existants (filled<5)
        overflow_groups = {}  # sid -> {village, insert_after, persons_list, start_num}
        new_village_list = [] # (quartier, persons_list) pour nouveaux villages
        phones_to_add  = []

        # Snapshot complet des slots
        slot_state = {}
        for row in con.execute("SELECT id,village,row_start,filled,last_data_row,arrond_norm FROM slots").fetchall():
            sid, village, row_start, filled, last_data_row, arrond_n = row
            slot_state[sid] = {"village": village, "row_start": row_start,
                               "filled": filled, "last_data_row": last_data_row}

        slot_fill_count  = {}  # sid -> filled count courant
        slot_lock        = {}  # (village_norm, arrond_ctx) -> sid (verrou de session)

        # ── Pré-charger les arronds en mémoire (évite 100s de requêtes SQL) ──
        arrond_cache = {}
        for row in con.execute(
                "SELECT arrond_norm, dept_norm, comm_norm, last_data_row, next_dept_row, hdr_row "
                "FROM arrond_index").fetchall():
            an, dn, cn, ldr, ndr, hr = row
            arrond_cache[an] = {"dept_norm": dn, "comm_norm": cn,
                                "last_data_row": ldr, "next_dept_row": ndr, "hdr_row": hr}

        # Pré-charger tous les slots en mémoire pour éviter les requêtes SQL répétées
        all_slots_by_norm = {}  # village_norm -> list of slot rows
        for row in con.execute(
                "SELECT id,village,row_start,filled,last_data_row,arrond_norm,dept_norm,comm_norm "
                "FROM slots").fetchall():
            sid, village, row_start, filled, last_data_row, arrond_n, dept_n, comm_n = row
            key = normalize(village)
            if key not in all_slots_by_norm:
                all_slots_by_norm[key] = []
            all_slots_by_norm[key].append(
                (sid, village, row_start, filled, last_data_row, arrond_n, dept_n, comm_n))

        def find_slot_fast(quartier, arrond_ctx=None, dept_ctx=None, comm_ctx=None):
            """Cherche un slot en mémoire (sans SQL) — beaucoup plus rapide."""
            key = normalize(quartier)
            candidates = all_slots_by_norm.get(key, [])
            if not candidates:
                if key not in _fuzzy_slot_cache:
                    matches = difflib.get_close_matches(key, _all_slot_keys, n=1, cutoff=0.75)
                    _fuzzy_slot_cache[key] = matches[0] if matches else None
                fuzzy_key = _fuzzy_slot_cache[key]
                if fuzzy_key:
                    candidates = all_slots_by_norm.get(fuzzy_key, [])
            if not candidates:
                return None
            # Niveau 1: dept + comm + arrond
            if dept_ctx and comm_ctx and arrond_ctx:
                for s in candidates:
                    if s[5] == arrond_ctx and s[6] == dept_ctx and s[7] == comm_ctx:
                        return s[:6]
            # Niveau 2: arrond seul
            if arrond_ctx:
                for s in candidates:
                    if s[5] == arrond_ctx:
                        return s[:6]
            # Niveau 3: global
            return candidates[0][:6]

        # Listes de clés pré-calculées pour fuzzy matching rapide
        _all_slot_keys   = list(all_slots_by_norm.keys())
        _all_arrond_keys = list(arrond_cache.keys())
        _fuzzy_slot_cache   = {}   # key → matched key ou None
        _fuzzy_arrond_cache = {}   # key → matched key ou None

        def find_arrond_fast(arrond_ctx, dept_ctx=None, comm_ctx=None):
            """Cherche un arrond en mémoire sans SQL."""
            if not arrond_ctx: return None
            if arrond_ctx in arrond_cache:
                a = arrond_cache[arrond_ctx]
                if not dept_ctx or a["dept_norm"] == dept_ctx:
                    return (arrond_ctx, a["last_data_row"], a["next_dept_row"], a["hdr_row"])
            if arrond_ctx not in _fuzzy_arrond_cache:
                matches = difflib.get_close_matches(arrond_ctx, _all_arrond_keys, n=1, cutoff=0.80)
                _fuzzy_arrond_cache[arrond_ctx] = matches[0] if matches else None
            fuzzy_key = _fuzzy_arrond_cache[arrond_ctx]
            if fuzzy_key:
                a = arrond_cache[fuzzy_key]
                return (fuzzy_key, a["last_data_row"], a["next_dept_row"], a["hdr_row"])
            return None

        for p in persons:
            tel = p["tel_clean"]
            if con.execute("SELECT 1 FROM phones WHERE phone=?", (tel,)).fetchone():
                duplicates += 1
                rejected.append({"row": p["row"], "quartier": p["quartier"], "nom": p["nom"],
                                 "reason": f"Doublon tel : {p['telephone']}"}); continue

            # Cherche le slot avec le contexte géographique complet
            arrond_ctx = p.get("arrond_ctx")
            dept_ctx   = p.get("dept_ctx")
            comm_ctx   = p.get("comm_ctx")

            # ── Verrou de session : même village → même slot ──────────────────
            lock_key = (normalize(p["quartier"]), arrond_ctx or "")
            if lock_key in slot_lock:
                # Réutiliser le slot déjà choisi pour ce village dans cette session
                locked_sid = slot_lock[lock_key]
                slot = next(
                    (s[:6] for s in all_slots_by_norm.get(normalize(p["quartier"]), [])
                     if s[0] == locked_sid),
                    None
                )
            else:
                slot = find_slot_fast(p["quartier"], arrond_ctx=arrond_ctx,
                                      dept_ctx=dept_ctx, comm_ctx=comm_ctx)
                if slot:
                    slot_lock[lock_key] = slot[0]  # verrouiller ce sid

            if slot:
                sid, village, row_start, filled, last_data_row, arrond_n = slot
                current_filled = slot_fill_count.get(sid, filled)

                if current_filled < 5:
                    target_row = row_start + current_filled
                    direct_writes.append((target_row, current_filled + 1, village, p))
                else:
                    if sid not in overflow_groups:
                        # Utiliser last_data_row depuis la DB (pas le cache slot_state)
                        # pour tenir compte des overflows des sessions précédentes
                        actual_last = con.execute(
                            "SELECT last_data_row FROM slots WHERE id=?", (sid,)
                        ).fetchone()[0]
                        overflow_groups[sid] = {
                            "village":      village,
                            "insert_after": actual_last,
                            "start_num":    current_filled + 1,
                            "persons":      []
                        }
                    overflow_groups[sid]["persons"].append(p)
                    overflow += 1

                slot_fill_count[sid] = current_filled + 1
                phones_to_add.append(tel)
                inserted += 1

            else:
                # ── Nouveau village → insérer à la fin de son arrondissement ──
                insert_after = _mere_max_row   # fallback = fin de fichier (pré-calculé)
                arrond_ctx = p.get("arrond_ctx")
                dept_ctx   = p.get("dept_ctx")
                comm_ctx   = p.get("comm_ctx")

                if arrond_ctx:
                    # Cherche l'arrond en mémoire (sans SQL — beaucoup plus rapide)
                    arrond_info = find_arrond_fast(arrond_ctx, dept_ctx, comm_ctx)
                    if arrond_info:
                        _, last_data, next_dept, hdr_row = arrond_info
                        insert_after = last_data  # juste après le dernier slot de l'arrond

                key = normalize(p["quartier"])
                existing = next((g for g in new_village_list if normalize(g["quartier"]) == key
                                and g.get("arrond_ctx") == arrond_ctx), None)
                if existing:
                    existing["persons"].append(p)
                else:
                    new_village_list.append({
                        "quartier": p["quartier"],
                        "insert_after": insert_after,
                        "arrond_ctx": arrond_ctx,
                        "persons": [p]
                    })
                new_quartier += 1
                phones_to_add.append(tel)
                inserted += 1        # ── Phase 2 : Trier les insertions par position DÉCROISSANTE ─────────
        sorted_overflows = sorted(overflow_groups.values(), key=lambda g: g["insert_after"], reverse=True)

        # ── Phase 3 : Styles pré-créés une seule fois (openpyxl est lent si recréés) ───
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side as Sd
        _thin    = Sd(style="thin", color="CCCCCC")
        _bdr     = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
        _fill    = PatternFill("solid", start_color=C_NEW)
        _font    = Font(name="Arial", size=10)
        _align_c = Alignment(horizontal="center", vertical="center")
        _align_v = Alignment(vertical="center")

        def style_inserted_row(ws, rn):
            ws.row_dimensions[rn].height = 15
            for c in range(1, NCOLS+1):
                cell       = ws.cell(row=rn, column=c)
                cell.fill  = _fill
                cell.font  = _font
                cell.border = _bdr
                cell.alignment = _align_c if c == 1 else _align_v

        def write_to_row(ws, rn, num, village, p):
            ws.cell(row=rn, column=1).value = num
            ws.cell(row=rn, column=2).value = village
            ws.cell(row=rn, column=3).value = p["nom"]
            ws.cell(row=rn, column=4).value = p["prenom"]
            ws.cell(row=rn, column=5).value = p.get("partis","")
            ws.cell(row=rn, column=6).value = p.get("profession","")
            ws.cell(row=rn, column=7).value = p.get("date_naissance","")
            ws.cell(row=rn, column=8).value = p.get("lieu_naissance","")
            ws.cell(row=rn, column=9).value = p["telephone"]

        # ── Phase 4 : Insertions groupées bottom-to-top (overflow) ───────────
        for group in sorted_overflows:
            count        = len(group["persons"])
            insert_at    = group["insert_after"] + 1  # insérer juste après dernière ligne
            village      = group["village"]
            start_num    = group["start_num"]

            # UNE SEULE insertion de N lignes → bloc contigu garanti
            ws_out.insert_rows(insert_at, amount=count)

            for i, p in enumerate(group["persons"]):
                rn = insert_at + i
                write_to_row(ws_out, rn, start_num + i, village, p)        # ── Phase 5 : Nouveaux villages GROUPÉS par position ─────────────────
        # Grouper les villages par insert_after → une seule insert_rows par groupe
        from collections import defaultdict as _dd
        _pos_groups = _dd(list)
        for nv in new_village_list:
            _pos_groups[nv["insert_after"]].append(nv)

        _nv_inserts = []  # (insert_at, total_count) pour adjust_row phase 6
        # Traiter du bas vers le haut (positions DESC) pour éviter le décalage
        for _pos in sorted(_pos_groups.keys(), reverse=True):
            _villages_at_pos = _pos_groups[_pos]
            _insert_at = _pos + 1
            _total = sum(len(nv["persons"]) for nv in _villages_at_pos)
            # UNE SEULE insert_rows pour toutes les villages de ce groupe
            ws_out.insert_rows(_insert_at, amount=_total)
            _nv_inserts.append((_insert_at, _total))
            _cur = _insert_at
            for nv in _villages_at_pos:
                for i, p in enumerate(nv["persons"]):
                    write_to_row(ws_out, _cur + i, i + 1, nv["quartier"], p)
                _cur += len(nv["persons"])        # ── Phase 6 : Écriture directe sur slots existants ───────────────────
        # IMPORTANT : recalculer les positions après les insert_rows
        # Chaque insert_rows(pos, n) décale de +n toutes les lignes >= pos
        # On accumule les décalages pour corriger les positions directes

        # Reconstruire la liste des décalages provoqués par les inserts (bottom-to-top → ordre inverse pour recalcul)
        # Les inserts ont été faits du bas vers le haut, donc recalculer du haut vers le bas
        all_inserts = []  # (insert_at, count)
        for group in sorted_overflows:
            all_inserts.append((group["insert_after"] + 1, len(group["persons"])))

        def adjust_row(original_row):
            """Recalcule la position après tous les insert_rows (overflow + nouveaux villages)."""
            adjusted = original_row
            for insert_at, count in all_inserts:
                if insert_at <= adjusted:
                    adjusted += count
            # Aussi ajuster pour les insertions de nouveaux villages
            for insert_at, count in _nv_inserts:
                if insert_at <= adjusted:
                    adjusted += count
            return adjusted

        for target_row, num, village, p in direct_writes:
            adjusted_row = adjust_row(target_row)
            ws_out.cell(row=adjusted_row, column=3).value = p["nom"]
            ws_out.cell(row=adjusted_row, column=4).value = p["prenom"]
            ws_out.cell(row=adjusted_row, column=5).value = p.get("partis","")
            ws_out.cell(row=adjusted_row, column=6).value = p.get("profession","")
            ws_out.cell(row=adjusted_row, column=7).value = p.get("date_naissance","")
            ws_out.cell(row=adjusted_row, column=8).value = p.get("lieu_naissance","")
            ws_out.cell(row=adjusted_row, column=9).value = p["telephone"]

        # ── Phase 7 : Téléphones (pas d'invalidation ici — on le fait après) ──
        for tel in phones_to_add:
            con.execute("INSERT OR IGNORE INTO phones VALUES(?)", (tel,))

        con.commit()
        wb_out.save(MERE_PATH)
    log = load_log()
    log["total_inserted"] = log.get("total_inserted", 0) + inserted
    if "users" not in log: log["users"] = {}
    # Per-user stats
    uname = username or "inconnu"
    if uname not in log["users"]:
        log["users"][uname] = {"total_inserted": 0, "sessions": []}
    log["users"][uname]["total_inserted"] = log["users"][uname].get("total_inserted", 0) + inserted
    log["users"][uname]["sessions"].append({
        "file": filename, "date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "inserted": inserted, "rejected": len(rejected),
        "duplicates": duplicates, "overflow": overflow, "new_quartiers": new_quartier,
    })
    # Global sessions
    log["sessions"].append({
        "file": filename, "date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "inserted": inserted, "rejected": len(rejected),
        "duplicates": duplicates, "overflow": overflow, "new_quartiers": new_quartier,
        "user": uname,
    })
    save_log(log)
    # Invalider + reconstruire l'index UNE seule fois si nécessaire
    if inserted > 0 or new_quartier > 0:
        invalidate_db()
    con_final = get_db()
    total_personnes = con_final.execute("SELECT SUM(filled) FROM slots").fetchone()[0] or 0
    return {
        "filename": filename, "inserted": inserted, "overflow": overflow,
        "new_quartiers": new_quartier, "duplicates": duplicates,
        "rejected_count": len(rejected), "rejected": rejected[:50],
        "total_mere": _mere_max_row, "total_personnes": total_personnes,
        "total_cumule": log["total_inserted"],
    }


# ── Auth Routes ───────────────────────────────────────────────────────────────

@app.route("/login", methods=["GET"])
def login_page():
    if is_logged_in():
        return redirect(url_for("index"))
    return render_template("login.html")

@app.route("/login", methods=["POST"])
def login():
    data = request.get_json() or request.form
    username = str(data.get("username","")).strip()
    password = str(data.get("password","")).strip()

    # Check admin
    if password == ADMIN_PASSWORD and (username.lower() in ["admin","administrateur"] or password == ADMIN_PASSWORD):
        if username.lower() in ["admin","administrateur"]:
            session["role"] = "admin"
            session["username"] = "admin"
            return jsonify({"ok": True, "role": "admin"})

    # Check regular user
    if check_password(username, password):
        session["username"] = username.lower().strip()
        session["display"]  = get_user(username)["display"]
        session["role"]     = "user"
        return jsonify({"ok": True, "role": "user"})

    return jsonify({"ok": False, "error": "Identifiants incorrects"}), 401

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login_page"))

# ── Main App Routes ───────────────────────────────────────────────────────────

@app.route("/")
@require_login
def index():
    return render_template("index.html",
        username=session.get("display", session.get("username","")),
        is_admin=is_admin())

@app.route("/upload", methods=["POST"])
@require_login
def upload():
    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "Aucun fichier recu"}), 400
    uname = current_username()
    results = []
    for f in files:
        if not f.filename.lower().endswith((".xlsx", ".xls")):
            results.append({"filename": f.filename, "error": "Format non supporte"}); continue

        # ── Nom unique par utilisateur + timestamp pour éviter collisions ──
        ext       = ".xls" if f.filename.lower().endswith(".xls") else ".xlsx"
        unique_id = f"{uname}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}{ext}"
        filepath  = os.path.join(UPLOAD_DIR, unique_id)

        f.save(filepath)
        try:
            result = integrate_file(filepath, f.filename, username=uname)
            results.append(result)
        except Exception as e:
            results.append({"filename": f.filename, "error": str(e)})
        finally:
            try: os.remove(filepath)
            except: pass
    return jsonify({"results": results, "log": load_log()})

@app.route("/download")
@require_admin
def download():
    return send_file(MERE_PATH, as_attachment=True,
                     download_name="Villages_Quartiers_Benin_Final.xlsx")

@app.route("/stats")
@require_login
def stats():
    log = load_log()
    con = get_db()
    total_personnes = con.execute("SELECT COUNT(*) FROM phones").fetchone()[0] or 0
    wb = load_workbook(MERE_PATH, read_only=True)
    ws = wb.active
    total_lignes = ws.max_row
    wb.close()

    uname = current_username()
    user_stats = log.get("users", {}).get(uname, {})
    user_sessions = user_stats.get("sessions", [])[-10:]
    user_total = user_stats.get("total_inserted", 0)

    return jsonify({
        "total_lignes_mere":   total_lignes,
        "total_personnes":     total_personnes,
        "total_insere_cumule": log.get("total_inserted", 0),
        "nb_sessions":         len(log.get("sessions", [])),
        "sessions":            log.get("sessions", [])[-10:],
        # User-specific
        "user_sessions":       user_sessions,
        "user_total_inserted": user_total,
        "user_nb_sessions":    len(user_sessions),
        "is_admin":            is_admin(),
    })

@app.route("/admin-panel")
@require_admin
def admin_panel_page():
    return render_template("admin.html")


@app.route("/admin/stats")
@require_admin
def admin_stats():
    """Stats complètes par utilisateur pour l'admin."""
    log = load_log()
    users_data = load_users()
    con = get_db()
    total_personnes = con.execute("SELECT SUM(filled) FROM slots").fetchone()[0] or 0

    user_summary = []
    for uname, udata in log.get("users", {}).items():
        display = users_data.get(uname, {}).get("display", uname) if uname != "__admin__" else "Admin"
        user_summary.append({
            "username":       uname,
            "display":        display,
            "total_inserted": udata.get("total_inserted", 0),
            "nb_sessions":    len(udata.get("sessions", [])),
            "last_session":   udata["sessions"][-1]["date"] if udata.get("sessions") else "—",
        })
    user_summary.sort(key=lambda x: x["total_inserted"], reverse=True)

    wb = load_workbook(MERE_PATH, read_only=True)
    ws = wb.active
    total_lignes = ws.max_row
    wb.close()

    return jsonify({
        "total_lignes_mere":   total_lignes,
        "total_personnes":     total_personnes,
        "total_insere_cumule": log.get("total_inserted", 0),
        "nb_sessions":         len(log.get("sessions", [])),
        "sessions":            log.get("sessions", [])[-20:],
        "users":               user_summary,
        "all_users":           list(users_data.keys()),
    })

@app.route("/admin/users", methods=["POST"])
@require_admin
def create_user_route():
    data = request.get_json() or {}
    username = str(data.get("username","")).strip()
    password = str(data.get("password","")).strip()
    if not username or not password:
        return jsonify({"ok": False, "error": "Prénom et mot de passe requis"}), 400
    ok, msg = create_user(username, password)
    return jsonify({"ok": ok, "message": msg})

@app.route("/admin/users/<username>", methods=["DELETE"])
@require_admin
def delete_user_route(username):
    users = load_users()
    key = username.lower().strip()
    if key not in users:
        return jsonify({"ok": False, "error": "Utilisateur introuvable"}), 404
    del users[key]
    save_users(users)
    return jsonify({"ok": True, "message": f"{username} supprimé"})

@app.route("/reset", methods=["POST"])
@require_admin
def reset():
    with write_lock:
        wb = load_workbook(MERE_PATH)
        ws = wb.active
        count = 0
        ORIGINAL_MAX = 28088
        for r in range(1, min(ORIGINAL_MAX+1, ws.max_row+1)):
            if str(ws.cell(row=r, column=1)) in ["1","2","3","4","5"]:
                for c in range(3, 10):
                    ws.cell(row=r, column=c).value = None
                count += 1
        if ws.max_row > ORIGINAL_MAX:
            ws.delete_rows(ORIGINAL_MAX+1, ws.max_row - ORIGINAL_MAX)
        wb.save(MERE_PATH)
        save_log({"sessions": [], "total_inserted": 0, "users": {}})
        invalidate_db()
    return jsonify({"message": f"{count} lignes reinitalisees"})

print("\n" + "="*55)
print("  \U0001f680 Villages & Quartiers Benin \u2014 Systeme v4")
print("  \u279c  http://localhost:5000")
print("="*55)

if not os.path.exists(MERE_PATH):
    print(f"\u274c Fichier mere introuvable : {MERE_PATH}")
else:
    print(f"\u2705 Fichier mere trouve : {MERE_PATH}")
    try:
        with app.app_context():
            get_db()
    except Exception as e:
        print(f"\u26a0\ufe0f  Index non construit : {e}")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
